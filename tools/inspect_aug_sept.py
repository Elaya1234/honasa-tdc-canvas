import os
import openpyxl

aug_path = os.path.join(r"C:\Users\elaya.n\Downloads", "Aug'26 SS ( HA) IM analysis .xlsx")
print(f"Path exists: {os.path.exists(aug_path)}")
if os.path.exists(aug_path):
    wb = openpyxl.load_workbook(aug_path, data_only=True)
    ws = wb['Sheet1']
    print(f"August dimensions: max_row={ws.max_row}, max_col={ws.max_column}")
    for r in range(16, 33):
        row_vals = [(c, ws.cell(r, c).value) for c in range(1, 15) if ws.cell(r, c).value is not None]
        if row_vals:
            print(f"Row {r}: {row_vals}")

sept_analysis_path = os.path.join(r"C:\Users\elaya.n\Documents\HONASA TDC\IM", "Sept'26 FW & SS IM analysis.xlsx")
if os.path.exists(sept_analysis_path):
    wb = openpyxl.load_workbook(sept_analysis_path, data_only=True)
    print("\nSept sheets:", wb.sheetnames)
    for s in wb.sheetnames:
        ws = wb[s]
        print(f"Sheet {s}: max_row={ws.max_row}, max_col={ws.max_column}")
