import os
import unittest
import openpyxl
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from tdc_im_agent.orchestrator import TDCOrchestrator
from tdc_im_agent.config import OUTPUT_DIR

class TestGoldenSeptemberDataset(unittest.TestCase):

    def setUp(self):
        self.input_file = os.path.join(BASE_DIR, "IM", "Sept'26 FW & SS.xlsx")
        self.assertTrue(os.path.exists(self.input_file), f"Golden test file missing: {self.input_file}")
        self.orchestrator = TDCOrchestrator()

    def test_inspect_and_recommend(self):
        """Verifies the agent's observation and inference without asking questions."""
        rec = self.orchestrator.inspect_and_recommend(self.input_file)
        
        self.assertEqual(rec["period_label"], "September 2026")
        self.assertEqual(rec["period_type"], "Monthly")
        self.assertEqual(rec["total_reels"], 46)
        self.assertGreaterEqual(rec["total_creators"], 40)
        self.assertLessEqual(rec["total_creators"], 46)
        self.assertEqual(rec["recommended_output"], "SKU-wise + Consolidated")
        self.assertIn("Language", rec["recommended_dimensions"])
        self.assertIn("Gender", rec["recommended_dimensions"])
        self.assertIn("Content Bucket", rec["recommended_dimensions"])
        self.assertIn("Content Category", rec["recommended_dimensions"])

    def test_end_to_end_analysis_and_reconciliation(self):
        """Executes full end-to-end analysis on the 46 Reels September dataset."""
        result = self.orchestrator.execute_analysis(
            file_path=self.input_file,
            download_media=False,
            reprocess_all=False
        )

        # 1. Output files & counts
        self.assertEqual(result["total_reels"], 46)
        self.assertEqual(result["successfully_processed"], 46)
        self.assertIn(result["verified_count"], [43, 44])
        self.assertEqual(result["no_speech_count"], 1)
        self.assertIn(result["caption_only_count"], [1, 2])

        # 2. QA & Reconciliation
        self.assertEqual(result["qa_status"], "PASSED")
        self.assertEqual(result["reconciliation_status"], "PASSED")
        self.assertEqual(result["completion_status"], "PARTIAL") # 2 records flagged as no speech / caption only

        # 3. Workbook physical inspection
        out_path = result["output_path"]
        self.assertTrue(os.path.exists(out_path))

        wb = openpyxl.load_workbook(out_path, data_only=False)
        expected_sheets = [
            "Executive Summary",
            "Consolidated All SKUs",
            "1% Hyalu Sunscreen Aqua Gel",
            "2% Sali-Cinamide Face Wash",
            "1% Kojic Acid Face Wash",
            "Benzoyl Peroxide Face Wash",
            "Data Quality & Verification",
            "Processing Log"
        ]
        for s in expected_sheets:
            self.assertIn(s, wb.sheetnames, f"Expected sheet '{s}' not found in generated workbook!")

        # 4. Verify Consolidated data row count
        ws_cons = wb["Consolidated All SKUs"]
        cons_rows = 0
        for r in range(2, 50):
            if ws_cons.cell(r, 1).value in ["Live", "live"]:
                cons_rows += 1
        self.assertEqual(cons_rows, 46, f"Consolidated sheet has {cons_rows} data rows, expected 46")

        # 5. Verify Individual SKU counts matching raw file
        ws_ha = wb["1% Hyalu Sunscreen Aqua Gel"]
        ha_cnt = sum(1 for r in range(2, 30) if ws_ha.cell(r, 1).value in ["Live", "live"])
        self.assertEqual(ha_cnt, 14, f"1% Hyalu Sunscreen Aqua Gel has {ha_cnt} rows, expected 14")

        ws_sali = wb["2% Sali-Cinamide Face Wash"]
        sali_cnt = sum(1 for r in range(2, 30) if ws_sali.cell(r, 1).value in ["Live", "live"])
        self.assertEqual(sali_cnt, 15, f"2% Sali-Cinamide Face Wash has {sali_cnt} rows, expected 15")

        ws_kojic = wb["1% Kojic Acid Face Wash"]
        kojic_cnt = sum(1 for r in range(2, 30) if ws_kojic.cell(r, 1).value in ["Live", "live"])
        self.assertEqual(kojic_cnt, 13, f"1% Kojic Acid Face Wash has {kojic_cnt} rows, expected 13")

        ws_bpo = wb["Benzoyl Peroxide Face Wash"]
        bpo_cnt = sum(1 for r in range(2, 30) if ws_bpo.cell(r, 1).value in ["Live", "live"])
        self.assertEqual(bpo_cnt, 4, f"Benzoyl Peroxide Face Wash has {bpo_cnt} rows, expected 4")

        # 14 + 15 + 13 + 4 = 46!
        self.assertEqual(ha_cnt + sali_cnt + kojic_cnt + bpo_cnt, 46)

        # 6. Verify Dynamic Formulas in Summary tables
        # Check summary header row in Consolidated sheet
        self.assertEqual(ws_cons.cell(50, 1).value, "Bucket")
        self.assertEqual(ws_cons.cell(50, 8).value, "Language")
        self.assertTrue(str(ws_cons.cell(51, 2).value).startswith("=SUMIF("))
        self.assertTrue(str(ws_cons.cell(51, 4).value).startswith("=COUNTIF("))
        self.assertTrue(str(ws_cons.cell(51, 5).value).startswith("=IF("))

        # Check Category & Gender summary table headers
        found_cat = False
        for r in range(52, 75):
            if ws_cons.cell(r, 1).value == "Category" and ws_cons.cell(r, 8).value == "Gender":
                found_cat = True
                break
        self.assertTrue(found_cat, "Category and Gender summary table not found in expected position")

        wb.close()

if __name__ == "__main__":
    unittest.main()
