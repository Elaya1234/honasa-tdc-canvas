import os
import sys
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from tdc_im_agent.cache_manager import CacheManager
from tdc_im_agent.models import CanonicalRow, ProcessingState, ScriptStatus
from tdc_im_agent.reel_discovery import extract_shortcode

def seed_september_cache():
    sept_analysis_path = os.path.join(BASE_DIR, "IM", "Sept'26 FW & SS IM analysis.xlsx")
    if not os.path.exists(sept_analysis_path):
        print(f"File not found: {sept_analysis_path}")
        return

    wb = openpyxl.load_workbook(sept_analysis_path, data_only=True)
    ws = wb["Consolidated All SKUs"]
    cache = CacheManager()

    count = 0
    for r in range(2, 48): # 46 rows
        link = str(ws.cell(r, 29).value or "").strip() # Col 29 (AC) is Live Link
        sc = extract_shortcode(link)
        if not sc:
            continue

        hook = str(ws.cell(r, 39).value or "").strip()
        script = str(ws.cell(r, 40).value or "").strip()
        status_txt = str(ws.cell(r, 41).value or "").strip()
        scraped_v = ws.cell(r, 42).value
        scraped_l = ws.cell(r, 43).value
        scraped_c = ws.cell(r, 44).value
        concept = str(ws.cell(r, 37).value or "").strip()
        bucket = str(ws.cell(r, 38).value or "").strip()
        lang = str(ws.cell(r, 17).value or "").strip()
        gender = str(ws.cell(r, 14).value or "").strip()

        row = CanonicalRow(
            source_row=r,
            creator=str(ws.cell(r, 3).value or "").strip(),
            live_link=link,
            shortcode=sc,
            spoken_hook=hook,
            verified_whole_script=script,
            script_scraping_status=status_txt or ScriptStatus.TRANSCRIBED.value,
            scraped_views=int(scraped_v) if scraped_v and str(scraped_v).isdigit() else None,
            scraped_likes=int(scraped_l) if scraped_l and str(scraped_l).isdigit() else None,
            scraped_comments=int(scraped_c) if scraped_c and str(scraped_c).isdigit() else None,
            concept_bucketing=concept,
            bucket=bucket or "Problem solution",
            language=lang or "English",
            gender=gender or "Female",
            processing_state=ProcessingState.SUCCESS if "Transcribed" in status_txt else ProcessingState.PARTIAL
        )

        cache.save_record(row)
        count += 1

    print(f"Successfully seeded {count} historical Reel records into persistent cache!")

if __name__ == "__main__":
    seed_september_cache()
