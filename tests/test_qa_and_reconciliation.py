import os
import unittest
import tempfile
import sys
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from tdc_im_agent.models import CanonicalRow, ScriptStatus, ProcessingState
from tdc_im_agent.report_generator import TDCReportGenerator
from tdc_im_agent.qa_validator import QAValidator
from tdc_im_agent.analytics_engine import AnalyticsEngine

class TestQAAndReconciliation(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def test_reconciliation_pass_on_valid_workbook(self):
        rows = [
            CanonicalRow(
                source_row=2,
                creator="Maniisha",
                product="2% Sali-Cinamide Anti-Acne Face Wash",
                sku_code="8904417319637",
                total_spends=25000.0,
                views_brand_collab=500000,
                likes=12000,
                comments=450,
                live_link="https://www.instagram.com/reel/C8q_sample1/",
                shortcode="C8q_sample1",
                spoken_hook="Acne marks are gone",
                verified_whole_script="Acne marks are gone with Sali-Cinamide...",
                script_scraping_status=ScriptStatus.TRANSCRIBED.value,
                processing_state=ProcessingState.SUCCESS
            ),
            CanonicalRow(
                source_row=3,
                creator="Shubham",
                product="1% Kojic Acid Face Wash",
                sku_code="8904417306248",
                total_spends=35000.0,
                views_brand_collab=700000,
                likes=18000,
                comments=600,
                live_link="https://www.instagram.com/reel/DLUO_sample2/",
                shortcode="DLUO_sample2",
                spoken_hook="Tan remove karo",
                verified_whole_script="Tan remove karo with Kojic Acid...",
                script_scraping_status=ScriptStatus.TRANSCRIBED.value,
                processing_state=ProcessingState.SUCCESS
            )
        ]

        exec_metrics = AnalyticsEngine.compute_executive_metrics(rows)
        test_out = os.path.join(self.temp_dir, "test_valid_report.xlsx")

        gen = TDCReportGenerator(period_label="September 2026")
        gen.build_full_workbook(
            rows=rows,
            dq_records=[],
            log_records=[],
            exec_summary_metrics=exec_metrics,
            output_path=test_out
        )

        qa_res = QAValidator.validate_workbook(test_out, rows)

        self.assertEqual(qa_res.qa_status, "PASSED")
        self.assertEqual(qa_res.reconciliation_status, "PASSED")
        self.assertEqual(qa_res.final_completion_state, "COMPLETE")
        self.assertTrue(qa_res.reels_reconciled)
        self.assertTrue(qa_res.spend_reconciled)
        self.assertTrue(qa_res.views_reconciled)
        self.assertEqual(len(qa_res.broken_formulas), 0)

    def test_broken_formula_detection(self):
        """Tests that QAValidator catches broken formula tokens like #REF! and #DIV/0!"""
        broken_file = os.path.join(self.temp_dir, "broken_formulas.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Consolidated All SKUs"
        ws.cell(1, 1, "Live")
        ws.cell(2, 1, "Live")
        ws.cell(2, 5, "=SUM(#REF!)") # Corrupted formula!
        
        # Add required sheets
        wb.create_sheet("Executive Summary")
        wb.create_sheet("Data Quality & Verification")
        wb.create_sheet("Processing Log")
        wb.save(broken_file)

        dummy_rows = [CanonicalRow(source_row=2, live_link="https://www.instagram.com/reel/abc/")]
        qa_res = QAValidator.validate_workbook(broken_file, dummy_rows)

        self.assertEqual(qa_res.qa_status, "FAILED")
        self.assertGreater(len(qa_res.broken_formulas), 0)

if __name__ == "__main__":
    unittest.main()
