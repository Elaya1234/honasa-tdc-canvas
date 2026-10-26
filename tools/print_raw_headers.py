import os
import openpyxl
import re
import pandas as pd

p_raw = os.path.join(r"c:\Users\elaya.n\Documents\HONASA TDC\IM", "Sept'26 FW & SS.xlsx")
wb = openpyxl.load_workbook(p_raw, data_only=True)
ws = wb['Sheet1']
headers = [str(ws.cell(1, c).value or '').strip() for c in range(1, ws.max_column + 1)]
print("Raw headers in Sept'26 FW & SS.xlsx:")
for idx, h in enumerate(headers, 1):
    print(f"  Col {idx} ({openpyxl.utils.get_column_letter(idx)}): '{h}'")
