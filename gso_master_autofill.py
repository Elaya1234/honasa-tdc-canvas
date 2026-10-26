import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

import os
import re
import time
import json
import urllib.parse
import openpyxl
from curl_cffi import requests
from bs4 import BeautifulSoup
from google import genai
from google.genai import types
from pydantic import BaseModel, Field

# ======================================================================
# MARKETPLACE LINK AUTO-FILLER
# Autonomous AI-powered link finder and Excel auto-filler for Honasa GSO
# ======================================================================

# Load API key
env_path = r"c:\Users\elaya.n\Documents\HONASA TDC\.env"
alt_env = r"c:\Users\elaya.n\Documents\HONASA TDC\ig_reel_collector\.env"
api_key = os.environ.get("GEMINI_API_KEY")

for p in [env_path, alt_env]:
    if not api_key and os.path.exists(p):
        with open(p, "r", encoding="utf-8") as f:
            for line in f:
                if line.startswith("GEMINI_API_KEY="):
                    api_key = line.strip().split("=", 1)[1].strip("\"' ")
                    break

if not api_key:
    print("[WARN] GEMINI_API_KEY not found in environment or .env files!")

client = None
if api_key:
    try:
        client = genai.Client(api_key=api_key)
    except Exception as e:
        print(f"[!] Warning: genai.Client init error: {e}")

class BatchSelection(BaseModel):
    selected_url: str = Field(description="The exact URL from the candidates that matches the single product, or empty string if none match.")
    is_match: bool = Field(description="True if an exact single-unit match was found.")
    reason: str = Field(description="Brief reason explaining the selection or rejection.")

def clean_url(url, domain):
    try:
        if "amazon.in" in domain:
            m = re.search(r'/(?:dp|gp/product)/([A-Z0-9]{10})', url)
            if m:
                return f"https://www.amazon.in/dp/{m.group(1)}"
            return url.split("?")[0]
        elif "flipkart.com" in domain:
            m = re.search(r'(/p/itm[a-zA-Z0-9]+)', url)
            if m:
                return f"https://www.flipkart.com{m.group(1)}"
            return url.split("?")[0]
        elif "nykaa.com" in domain:
            m = re.search(r'(/p/\d+)', url)
            if m:
                return url.split("?")[0]
            return url.split("?")[0]
        elif "zepto" in domain:
            m = re.search(r'(/pn/[^/]+/pvid/[^/?]+)', url)
            if m:
                return f"https://www.zepto.com{m.group(1)}"
            return url.split("?")[0]
        elif "blinkit" in domain:
            m = re.search(r'(/prn/[^/]+/prid/\d+)', url)
            if m:
                return f"https://blinkit.com{m.group(1)}"
            return url.split("?")[0]
        else:
            return url.split("?")[0]
    except:
        return url

def search_candidates_ddg(query, domain=None):
    encoded = urllib.parse.quote_plus(query)
    url = f"https://html.duckduckgo.com/html/?q={encoded}"
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"}
    results = []
    try:
        r = requests.get(url, impersonate="chrome120", headers=headers, timeout=8)
        soup = BeautifulSoup(r.text, "html.parser")
        for div in soup.find_all("div", class_="result"):
            a_link = div.find("a", class_="result__snippet")
            if not a_link:
                continue
            href = a_link.get("href", "")
            if "uddg=" in href:
                actual = urllib.parse.unquote(href.split("uddg=")[1].split("&")[0])
                if "y.js" not in actual and "ad_domain" not in actual and "bing.com" not in actual:
                    if domain is None or domain in actual:
                        title_text = div.find("h2").get_text(strip=True) if div.find("h2") else ""
                        snippet_text = a_link.get_text(strip=True) if a_link else ""
                        cleaned = clean_url(actual, domain or "")
                        if cleaned not in [res["url"] for res in results]:
                            results.append({
                                "url": cleaned,
                                "search_title": title_text,
                                "snippet": snippet_text
                            })
    except Exception as e:
        pass
    return results[:5]

def enrich_candidate(candidate):
    url = candidate["url"]
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"}
    try:
        r = requests.get(url, impersonate="chrome120", headers=headers, timeout=6)
        if r.status_code == 200:
            soup = BeautifulSoup(r.text, "html.parser")
            az_title = soup.find(id="productTitle")
            if az_title:
                candidate["page_title"] = az_title.get_text(strip=True)
            elif soup.title:
                candidate["page_title"] = soup.title.string or ""
            
            meta_desc = soup.find("meta", attrs={"name": "description"}) or soup.find("meta", attrs={"property": "og:description"})
            if meta_desc:
                candidate["page_desc"] = meta_desc.get("content", "")
    except:
        pass
    return candidate

def evaluate_candidates_with_ai(target_brand, target_variant, platform_name, candidates):
    if not candidates:
        return None
        
    candidates_text = ""
    for idx, c in enumerate(candidates, 1):
        title = c.get("page_title") or c.get("search_title", "")
        desc = c.get("page_desc") or c.get("snippet", "")
        clean_title = title.encode('ascii', 'ignore').decode('ascii')
        clean_desc = desc.encode('ascii', 'ignore').decode('ascii')
        candidates_text += f"\nCandidate {idx}:\n  URL: {c['url']}\n  Title: {clean_title}\n  Description/Snippet: {clean_desc}\n"

    prompt = f"""
You are an expert strict e-commerce data verification agent for Indian retail marketplaces.

TARGET PRODUCT:
Brand: {target_brand}
Target Variant & Size: {target_variant} (SINGLE UNIT ONLY)
Platform: {platform_name}

CANDIDATES FOUND:{candidates_text}

VERIFICATION RULES:
1. Exact Brand & Core Variant: The brand and core variant must match (e.g. Aqua Gel vs Matte Gel vs Cica Matte). Ignore general retailer marketing fluff like 'SPF 50 PA++++', 'with Vitamin E', 'Broad Spectrum', 'For Glowing Skin'.
2. Single Unit Only: Must be a single unit matching target size (50g/50ml/30ml/100ml). Strictly REJECT multi-packs, 'Pack of 2', 'Pack of 3', combos, and bundles.
3. Size Conflict: If a candidate explicitly states a conflicting size (e.g. 80g, 125g, 30g, 100g) when 50g was requested, reject it. If no size conflict is present and it is the standard single unit, accept it.
4. Promotional Slogans: Site-wide promotions in title (e.g. '| Buy 1 Get 1 FREE' on brand D2C) represent storewide promo banners, NOT a bundle, unless the item title explicitly says 'Pack of 2' or 'Combo'.

Select the single best matching URL from the candidates, or return is_match=False if none match.
"""
    models_to_try = ["gemini-3.5-flash-lite", "gemini-2.5-flash", "gemini-flash-lite-latest", "gemini-2.0-flash"]
    for m in models_to_try:
        try:
            resp = client.models.generate_content(
                model=m,
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema=BatchSelection,
                    temperature=0.1
                )
            )
            data = json.loads(resp.text)
            return data
        except Exception as e:
            continue
    print(f"    [!] All Gemini model attempts failed for candidate evaluation")
    return None

D2C_DOMAINS = {
    "The Derma Co": "thedermaco.com",
    "Minimalist": "beminimalist.co",
    "Dot & Key": "dotandkey.com",
    "L'Oreal Paris": "lorealparis.co.in",
    "L'Oral Paris": "lorealparis.co.in",
    "Foxtale": "foxtale.in",
    "Cetaphil": "cetaphil.in",
    "The Ordinary": "theordinary.com",
    "Himalaya": "himalayawellness.in",
    "Simple": "simpleskincare.in"
}

MARKETPLACE_CONFIG = [
    {"col": 5, "name": "Brand D2C", "domain_fn": lambda b: D2C_DOMAINS.get(b, ""), "query_fn": lambda b, v: f"site:{D2C_DOMAINS.get(b, '')} {v}" if D2C_DOMAINS.get(b) else ""},
    {"col": 6, "name": "Amazon", "domain_fn": lambda b: "amazon.in", "query_fn": lambda b, v: f"Amazon India {b} {v}"},
    {"col": 7, "name": "Flipkart", "domain_fn": lambda b: "flipkart.com", "query_fn": lambda b, v: f"Flipkart {b} {v}"},
    {"col": 8, "name": "Nykaa", "domain_fn": lambda b: "nykaa.com", "query_fn": lambda b, v: f"Nykaa {b} {v}"},
    {"col": 9, "name": "Blinkit", "domain_fn": lambda b: "blinkit.com", "query_fn": lambda b, v: f"Blinkit {b} {v}"},
    {"col": 10, "name": "Zepto", "domain_fn": lambda b: "zepto.com", "query_fn": lambda b, v: f"Zepto {b} {v}"},
    {"col": 11, "name": "Instamart", "domain_fn": lambda b: "swiggy.com", "query_fn": lambda b, v: f"Swiggy Instamart {b} {v}"},
]

def safe_save(wb, path):
    try:
        wb.save(path)
        return True
    except PermissionError:
        alt_path = path.replace(".xlsx", "_temp.xlsx")
        print(f"  [!] File locked! Saving backup to: {alt_path}")
        wb.save(alt_path)
        return False

def process_product_marketplaces(brand, variant, ws, row_idx):
    print(f"\n=======================================================")
    print(f"Row {row_idx}: {brand} | {variant}")
    print(f"=======================================================")
    
    for mp in MARKETPLACE_CONFIG:
        col = mp["col"]
        name = mp["name"]
        
        # Check if already filled
        existing = ws.cell(row=row_idx, column=col).value
        if existing and str(existing).strip().startswith("http"):
            print(f"  [ALREADY FILLED] {name}: {str(existing)[:60]}")
            continue
            
        domain = mp["domain_fn"](brand)
        if not domain and name == "Brand D2C":
            continue
            
        query = mp["query_fn"](brand, variant)
        if not query:
            continue
            
        print(f"  Searching {name}...")
        candidates = search_candidates_ddg(query, domain)
        if not candidates:
            fallback_query = f"{name} {brand} {variant}"
            candidates = search_candidates_ddg(fallback_query, domain)
            
        if not candidates:
            print(f"    -> No candidates found on {name}.")
            continue
            
        print(f"    Found {len(candidates)} candidates. Quick-enriching...")
        for c in candidates[:3]:
            enrich_candidate(c)
            
        verdict = evaluate_candidates_with_ai(brand, variant, name, candidates)
        if verdict and verdict.get("is_match") and verdict.get("selected_url"):
            matched_url = verdict["selected_url"]
            print(f"    [AI APPROVED] {name}: {matched_url}")
            print(f"    Reason: {verdict.get('reason')}")
            ws.cell(row=row_idx, column=col).value = matched_url
        else:
            reason = verdict.get('reason') if verdict else 'No evaluation'
            print(f"    [AI REJECTED ALL] {name} -> {reason}")
            
        time.sleep(1.0)

def run_autofill(excel_path=None, start_row=72, end_row=88):
    if not excel_path:
        default_verified = r"c:\Users\elaya.n\Documents\HONASA TDC\GSO_LINKs_AI_Verified.xlsx"
        default_raw = r"c:\Users\elaya.n\Documents\HONASA TDC\GSO LINKs.xlsx"
        excel_path = default_verified if os.path.exists(default_verified) else default_raw
        
    print(f"\n=======================================================")
    print(f"   MARKETPLACE LINK AUTO-FILLER")
    print(f"   Target Workbook: {os.path.basename(excel_path)}")
    print(f"   Processing Rows: {start_row} to {end_row}")
    print(f"=======================================================\n")

    if not os.path.exists(excel_path):
        print(f"[!] Error: File not found: {excel_path}")
        return

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active if "Sheet1" not in wb.sheetnames else wb["Sheet1"]
    
    for r in range(start_row, end_row + 1):
        c2 = ws.cell(row=r, column=3).value
        c3 = ws.cell(row=r, column=4).value
        if not c2 or not c3:
            continue
        brand = str(c2).strip()
        variant = str(c3).strip()
        if len(variant) < 3 or not any(c.isalnum() for c in variant):
            continue
            
        process_product_marketplaces(brand, variant, ws, r)
        safe_save(wb, excel_path)
        print(f"Saved progress after Row {r}.")
        
    print("\n[COMPLETE] Marketplace Link Auto-Filler processing completed and saved!")

if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else None
    s_row = int(sys.argv[2]) if len(sys.argv) > 2 else 72
    e_row = int(sys.argv[3]) if len(sys.argv) > 3 else 88
    run_autofill(excel_path=target, start_row=s_row, end_row=e_row)
