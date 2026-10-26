import os
import re
import openpyxl
from typing import Dict, List, Set

COLUMN_ALIASES: Dict[str, List[str]] = {
    "live_status": ["live", "status", "video status", "post status", "asset status"],
    "sub_cat": ["sub cat", "subcat", "sub category", "subcategory", "category group"],
    "creator": ["creator", "influencer", "influencer name", "creator name", "artist"],
    "channel_standard_name": ["channel standard name", "standard name", "handle", "creator handle", "instagram handle"],
    "creator_ig_link": ["creator ig link", "creator link", "ig link", "profile link", "instagram profile", "creator profile"],
    "product": ["sept product", "aug product", "product", "sku name", "product name", "item name", "sku"],
    "sku_code": ["sku code", "barcode", "ean", "ean code", "sku_code"],
    "stage": ["stage", "format", "deliverable"],
    "epd_npd": ["epd/npd", "epd", "npd", "product stage"],
    "count": ["count", "reel count", "qty", "quantity"],
    "range_col": ["range", "brand range", "product range"],
    "asset": ["assest", "asset", "creative asset"],
    "agency": ["agency", "talent agency"],
    "gender": ["gender", "creator gender"],
    "follower_base": ["follower base", "followers", "follower count", "audience size"],
    "tier": ["tier", "influencer tier", "creator tier"],
    "language": ["language", "spoken language", "lang"],
    "genre": ["genre", "content genre"],
    "category": ["category", "creator category", "influencer category"],
    "er_planned": ["er", "planned er", "target er"],
    "average_views": ["average views", "avg views", "past avg views"],
    "planned_cpv": ["cpv", "planned cpv"],
    "cost": ["cost", "base cost", "creator cost"],
    "counter": ["counter", "counter cost"],
    "final_costs": ["final costs", "final cost", "agreed cost", "deal cost"],
    "commission": ["commission", "agency commission", "agency fee"],
    "total_spends": ["total spends", "total spend", "spends", "spend", "budget", "cost total", "investment"],
    "video_status": ["video status", "delivery status"],
    "live_link": ["live link", "reel link", "creative link", "reel url", "link", "post link", "instagram url", "url"],
    "live_date": ["live date", "posting date", "date", "post date", "publish date"],
    "views_organic": ["views day 0-5", "views day 0–5", "views day 0-5 (organic - creator handle)", "organic views"],
    "views_brand_collab": ["views day 10th", "views day 10", "views day 10th (organic - brand collab)", "views brand collab", "collab views"],
    "video_cpv": ["video cpv", "actual cpv", "realized cpv"],
    "likes": ["likes", "like count", "total likes"],
    "comments": ["comments", "comment count", "total comments"],
    "actual_er": ["er%", "actual er", "engagement rate"],
    "concept_bucketing": ["concept bucketing", "concept style", "style"],
    "bucket": ["bucket", "content bucket", "macro bucket", "concept"],
    "spoken_hook": ["spoken hook", "hook", "opening hook"],
    "verified_whole_script": ["verified whole script", "whole script", "script", "transcript"],
    "script_scraping_status": ["script scraping status", "script status", "scraping status"],
    "scraped_views": ["scraped views", "live views"],
    "scraped_likes": ["scraped likes", "live likes"],
    "scraped_comments": ["scraped comments", "live comments"],
    "verification_notes": ["verification notes", "notes", "audit notes"]
}

p_raw = os.path.join(r"c:\Users\elaya.n\Documents\HONASA TDC\IM", "Sept'26 FW & SS.xlsx")
wb = openpyxl.load_workbook(p_raw, data_only=True)
ws = wb['Sheet1']
raw_cols = [re.sub(r'[\r\n]+', ' ', str(ws.cell(1, c).value or '')).strip() for c in range(1, ws.max_column + 1)]

field_to_col: Dict[str, str] = {}
used_cols: Set[str] = set()

# Pass 1: Exact Matches
for canonical_field, aliases in COLUMN_ALIASES.items():
    for col_name in raw_cols:
        clean_col = col_name.lower().replace('–', '-').replace('—', '-').strip()
        for alias in aliases:
            if clean_col == alias.lower().strip():
                field_to_col[canonical_field] = col_name
                used_cols.add(col_name)
                break
        if canonical_field in field_to_col:
            break

# Pass 2: Specialized token matching
for canonical_field, aliases in COLUMN_ALIASES.items():
    if canonical_field in field_to_col:
        continue
    for col_name in raw_cols:
        if col_name in used_cols:
            continue
        clean_col = col_name.lower().replace('–', '-').replace('—', '-').strip()
        
        # Specific guards
        if canonical_field == "live_link" and any(x in clean_col for x in ["creator", "profile", "channel"]):
            continue
        if canonical_field == "views_brand_collab" and any(x in clean_col for x in ["average", "avg", "0-5", "0–5"]):
            continue
        if canonical_field == "views_organic" and not any(x in clean_col for x in ["0-5", "0–5", "organic"]):
            continue
        if canonical_field == "count" and clean_col == "counter":
            continue

        matched = False
        for alias in aliases:
            tokens_alias = set(re.findall(r'\w+', alias.lower()))
            tokens_col = set(re.findall(r'\w+', clean_col))
            if tokens_alias and tokens_alias.issubset(tokens_col):
                field_to_col[canonical_field] = col_name
                used_cols.add(col_name)
                matched = True
                break
        if matched:
            break

print("Pass 1 & Pass 2 Mappings:")
for k, v in sorted(field_to_col.items()):
    print(f"  {k:25} -> '{v}'")

print("\nKey validation:")
print("  live_link:", field_to_col.get("live_link"))
print("  creator_ig_link:", field_to_col.get("creator_ig_link"))
print("  views_brand_collab:", field_to_col.get("views_brand_collab"))
print("  views_organic:", field_to_col.get("views_organic"))
print("  count:", field_to_col.get("count"))
print("  counter:", field_to_col.get("counter"))
print("  total_spends:", field_to_col.get("total_spends"))
