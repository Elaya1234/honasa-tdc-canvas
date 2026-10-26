import io
import csv
from typing import List
from models import ReelProcessedRecord

COLUMNS = [
    "Reel URL",
    "Channel Standard Name",
    "Creator IG Link",
    "Current Views",
    "Likes",
    "Comments",
    "Language",
    "Whole script extraction",
    "Hook",
    "Genre / Category / Sub Cat",
    "Concept Bucketing",
    "Gender",
    "Status"
]

def records_to_csv_bytes(records: List[ReelProcessedRecord]) -> bytes:
    """Exports records to UTF-8-SIG CSV bytes suitable for Windows Excel."""
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow(COLUMNS)

    for r in records:
        writer.writerow([
            r.reel_url,
            r.channel_standard_name,
            r.creator_ig_link,
            r.current_views,
            r.likes,
            r.comments,
            r.language,
            r.whole_script,
            r.hook,
            r.genre_category_subcat,
            r.concept_bucketing,
            r.gender,
            r.status
        ])

    return output.getvalue().encode("utf-8-sig")


def records_to_excel_bytes(records: List[ReelProcessedRecord]) -> bytes:
    """
    Exports records to beautifully formatted Excel (.xlsx) bytes with openpyxl.
    Falls back to CSV bytes if openpyxl is not installed.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reel Analysis"
        ws.views.sheetView[0].showGridLines = True

        # Header Styling
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.append(COLUMNS)
        ws.row_dimensions[1].height = 28

        for col_num in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        # Body Styling
        body_font = Font(name="Segoe UI", size=10)
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )

        for row_idx, r in enumerate(records, start=2):
            ws.append([
                r.reel_url,
                r.channel_standard_name,
                r.creator_ig_link,
                r.current_views,
                r.likes,
                r.comments,
                r.language,
                r.whole_script,
                r.hook,
                r.genre_category_subcat,
                r.concept_bucketing,
                r.gender,
                r.status
            ])
            ws.row_dimensions[row_idx].height = 22

            is_alt = (row_idx % 2 == 1)
            for col_num in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = body_font
                cell.border = thin_border
                if is_alt:
                    cell.fill = alt_fill

                # Alignment and number formatting
                col_name = COLUMNS[col_num - 1]
                if col_name in ["Current Views", "Likes", "Comments"]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"
                elif col_name in ["Language", "Gender", "Status"]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                # Limit checking length for long script cells
                line_len = len(val.split("\n")[0]) if "\n" in val else len(val)
                max_len = max(max_len, min(line_len, 45))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    except ImportError:
        # Fallback to CSV
        return records_to_csv_bytes(records)
