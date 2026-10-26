import os
import openpyxl
import pandas as pd

def inspect_workbook(filepath):
    print("=" * 70)
    print(f"FILE: {filepath}")
    print("=" * 70)
    if not os.path.exists(filepath):
        print("ERROR: File does not exist!")
        return

    wb = openpyxl.load_workbook(filepath, data_only=False)
    print(f"Sheet names ({len(wb.sheetnames)}): {wb.sheetnames}")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n--- Sheet: '{sheet_name}' (max_row={ws.max_row}, max_col={ws.max_column}) ---")
        
        # Check merged ranges
        if ws.merged_cells.ranges:
            print(f"  Merged ranges count: {len(ws.merged_cells.ranges)} (e.g., {[str(r) for r in list(ws.merged_cells.ranges)[:5]]})")

        # Check frozen panes
        if ws.freeze_panes:
            print(f"  Freeze panes: {ws.freeze_panes}")

        # Find header rows and sample data
        sample_rows = []
        for r in range(1, min(12, ws.max_row + 1)):
            vals = [ws.cell(r, c).value for c in range(1, min(35, ws.max_column + 1))]
            if any(v is not None for v in vals):
                sample_rows.append((r, vals))

        print(f"  First non-empty rows:")
        for r, vals in sample_rows[:6]:
            non_empty = [(idx + 1, str(v)[:30]) for idx, v in enumerate(vals) if v is not None]
            print(f"    Row {r}: {non_empty[:8]}")

        # Check for formulas in the sheet
        formulas_found = []
        for r in range(1, min(ws.max_row + 1, 60)):
            for c in range(1, min(ws.max_column + 1, 40)):
                v = str(ws.cell(r, c).value or '')
                if v.startswith('='):
                    formulas_found.append((f"{openpyxl.utils.get_column_letter(c)}{r}", v))
                    if len(formulas_found) >= 8:
                        break
            if len(formulas_found) >= 8:
                break
        if formulas_found:
            print(f"  Sample formulas: {formulas_found}")

files = [
    r"C:\Users\elaya.n\Downloads\Aug'26 SS ( HA) IM analysis .xlsx",
    r"c:\Users\elaya.n\Documents\HONASA TDC\IM\Sept'26 FW & SS.xlsx",
    r"c:\Users\elaya.n\Documents\HONASA TDC\IM\Sept'26 FW & SS IM analysis.xlsx",
]

for f in files:
    inspect_workbook(f)
