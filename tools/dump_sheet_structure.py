import os
import openpyxl

def dump_sheet_structure(filepath, sheet_name):
    print(f"\n==========================================")
    print(f"FILE: {os.path.basename(filepath)} | SHEET: {sheet_name}")
    print(f"==========================================")
    wb = openpyxl.load_workbook(filepath, data_only=False)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found! Available: {wb.sheetnames}")
        return
    ws = wb[sheet_name]
    print(f"Dimensions: max_row={ws.max_row}, max_col={ws.max_column}")
    
    # Headers in Row 1
    headers = []
    for c in range(1, ws.max_column + 1):
        col_letter = openpyxl.utils.get_column_letter(c)
        v = ws.cell(1, c).value
        if v is not None:
            headers.append((c, col_letter, str(v).replace('\n', ' ')))
    print("\nRow 1 Headers:")
    for c, letter, name in headers:
        print(f"  Col {c} ({letter}): {name}")

    # Inspect summary tables (rows after data, or where 'Bucket', 'Summary', etc. occur)
    print("\nScanning rows for summary tables / blocks:")
    for r in range(1, ws.max_row + 1):
        c1 = str(ws.cell(r, 1).value or '')
        c2 = str(ws.cell(r, 2).value or '')
        c8 = str(ws.cell(r, 8).value or '')
        c15 = str(ws.cell(r, 15).value or '')
        if any(keyword in (c1 + c2 + c8 + c15).lower() for keyword in ['bucket', 'language', 'category', 'gender', 'summary', 'grand total', 'total']):
            row_items = [(openpyxl.utils.get_column_letter(c) + str(r), str(ws.cell(r, c).value)[:35]) 
                         for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]
            print(f"  Row {r}: {row_items[:10]}")

f_aug = r"C:\Users\elaya.n\Downloads\Aug'26 SS ( HA) IM analysis .xlsx"
dump_sheet_structure(f_aug, "Consolidated")
if "1% HA Sunscreen" in openpyxl.load_workbook(f_aug, read_only=True).sheetnames:
    dump_sheet_structure(f_aug, "1% HA Sunscreen")
else:
    dump_sheet_structure(f_aug, openpyxl.load_workbook(f_aug, read_only=True).sheetnames[0])

f_sept_analysis = r"c:\Users\elaya.n\Documents\HONASA TDC\IM\Sept'26 FW & SS IM analysis.xlsx"
wb_sept = openpyxl.load_workbook(f_sept_analysis, read_only=True)
print("Sept analysis sheetnames:", wb_sept.sheetnames)
dump_sheet_structure(f_sept_analysis, "Consolidated")
dump_sheet_structure(f_sept_analysis, "Benzoyl Peroxide Face Wash")

f_sept_raw = r"c:\Users\elaya.n\Documents\HONASA TDC\IM\Sept'26 FW & SS.xlsx"
wb_sept_raw = openpyxl.load_workbook(f_sept_raw, read_only=True)
print("Sept raw sheetnames:", wb_sept_raw.sheetnames)
dump_sheet_structure(f_sept_raw, wb_sept_raw.sheetnames[0])
